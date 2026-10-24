#!/usr/bin/env python3
"""Gera o relatório (PDF e DOCX) e os slides (PPTX e PDF) da AV1 a partir das fontes do repositório.

Entradas:
  docs/relatorio/RELATORIO_AV1.md      texto do relatório, com marcadores
                                        <!-- incluir: caminho -->  e  <!-- figura: caminho | legenda -->
  docs/evidencias/resultados/*.md|.json|.svg   gerados por EvidenciasAv1Test / scripts/av1_evidencias.ps1
Saídas:
  docs/relatorio/RELATORIO_AV1.pdf e .docx
  docs/presentations/apresentacao_I_AV1.pptx e .pdf

Requer: pandoc, python-playwright (Chromium), python-pptx e LibreOffice (soffice) no PATH; rsvg-convert é opcional
(melhora a figura de dispersão nos slides). Ícones dos slides: docs/presentations/icones, gerados por
scripts/gerar_icones_slides.js (Node, só para refazer os ícones).
Nenhum número é digitado aqui: os valores dos slides vêm dos .json de resultados.

Uso: python scripts/gerar_documentos_av1.py
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
DOCS = RAIZ / "docs"
EVID = DOCS / "evidencias"
RES = EVID / "resultados"
REL = DOCS / "relatorio"
APRES = DOCS / "presentations"
FONTE = RAIZ / "app/src/main/res/font/plus_jakarta_sans.ttf"

VIOLETA = "#5E4AE3"
VERDE = "#0D9F6E"
VERMELHO = "#E03E3E"
ESCURO = "#0D0F14"
CLARO = "#F4F5F9"

# Quem apresenta. Por bloco: aparece na capa. Por slide: aparece no rodapé. Vazio: só o nome do bloco.
APRESENTADORES = {
    "Fundamentação": "João Paulo e Juliano",
    "MVP ao vivo": "Juliano, Andre e João Paulo",
    "Resultados": "Andre",
}
APRESENTADOR_SLIDE = {
    2: "João Paulo Silva", 3: "Juliano Cesar Ferreira Ramos", 4: "Juliano Cesar Ferreira Ramos",
    5: "João Paulo Silva", 6: "João Paulo Silva", 7: "Juliano Cesar Ferreira Ramos", 8: "João Paulo Silva",
    9: "Juliano, Andre e João Paulo", 10: "Andre Bueno Ocanha", 11: "Andre Bueno Ocanha",
}


def rodape(bloco, extra="", numero=None):
    partes = [bloco, APRESENTADOR_SLIDE.get(numero) or APRESENTADORES.get(bloco, ""), extra]
    return " · ".join(x for x in partes if x)


def pendente(texto):
    return f"⟪Pendente: {texto}⟫"


def resolver_inclusoes(md: str) -> str:
    def incluir(m):
        partes = [p.strip() for p in m.group(1).split("|")]
        rel, modo = partes[0], (partes[1] if len(partes) > 1 else "")
        alvo = EVID / rel
        if not alvo.exists():
            return pendente(f"gerar `{rel}` com `scripts\\av1_evidencias.ps1` depois da coleta (docs/ENSAIO_AV1.md)")
        texto = alvo.read_text(encoding="utf-8")
        linhas = texto.splitlines()
        if linhas and linhas[0].startswith("# "):
            linhas = linhas[1:]
        saida = []
        dentro_codigo = False
        for l in linhas:
            if l.startswith("```"):
                dentro_codigo = not dentro_codigo
            if not dentro_codigo and l.startswith("#"):
                l = "###" + l.lstrip("#")  # rebaixa títulos internos
            saida.append(l)
        texto = "\n".join(saida).strip()
        # A saída bruta do Gradle é longa: no PDF fica só o resumo (o arquivo completo está no repositório).
        if rel.endswith("testes_unidade.md"):
            texto = re.sub(r"```text.*?```", "_Saída completa do Gradle em `docs/evidencias/resultados/testes_unidade.md`._",
                           texto, flags=re.S)
            if modo == "resumo":
                # No corpo do relatório: só o cabeçalho e os totais; a tabela por classe vai no apêndice.
                linhas = re.findall(r"^\|\s*(com\.[^|]+)\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|", texto, flags=re.M)
                tot = [sum(int(l[i]) for l in linhas) for i in range(1, 5)]
                cab = [l for l in texto.splitlines() if l.startswith("- ")]
                texto = "\n".join(cab) + (f"\n\n**Total: {tot[0]} testes em {len(linhas)} classes, {tot[1]} falhas, {tot[2]} erros, "
                                          f"{tot[3]} ignorados.** A tabela por classe está no Apêndice B.")
        return texto

    def figura(m):
        rel, legenda = [p.strip() for p in m.group(1).split("|", 1)]
        alvo = EVID / rel
        if not alvo.exists():
            # Capturas de tela podem vir como .png ou .jpg; aceita qualquer uma com o mesmo nome.
            alternativas = [alvo.with_suffix(ext) for ext in (".png", ".jpg", ".jpeg", ".svg")]
            alvo = next((a for a in alternativas if a.exists()), alvo)
        if not alvo.exists():
            origem = "gerada por EvidenciasAv1Test" if rel.startswith("resultados/") else "salvar a captura com este nome"
            return pendente(f"figura `docs/evidencias/{rel}` ({origem})")
        return f"![{legenda}]({alvo.as_posix()})"

    md = re.sub(r"<!--\s*incluir:\s*(.+?)\s*-->", incluir, md)
    md = re.sub(r"<!--\s*figura:\s*(.+?)\s*-->", figura, md)
    return md


CSS = """
@font-face { font-family: 'Plus Jakarta Sans'; src: url('FONTE') format('truetype'); font-weight: 200 800; }
@page { size: A4; margin: 18mm 16mm 18mm 16mm; }
body { font-family: 'Plus Jakarta Sans', 'Helvetica Neue', Arial, sans-serif; font-size: 10pt; line-height: 1.42; color: #0D0F14; }
header#title-block-header { border-bottom: 3px solid #5E4AE3; margin-bottom: 14pt; padding-bottom: 8pt; }
h1.title { font-size: 21pt; margin: 0; color: #5E4AE3; }
p.subtitle { font-size: 11pt; margin: 3pt 0 0 0; color: #444; }
p.date { font-size: 9pt; color: #666; margin: 2pt 0 0 0; }
h1 { font-size: 14pt; color: #5E4AE3; margin: 16pt 0 5pt 0; break-after: avoid; }
h2 { font-size: 11.5pt; margin: 12pt 0 4pt 0; break-after: avoid; }
h3 { font-size: 10.5pt; margin: 9pt 0 3pt 0; break-after: avoid; }
p { margin: 0 0 6pt 0; text-align: justify; }
code { font-family: 'DejaVu Sans Mono', Consolas, monospace; font-size: 8.6pt; background: #F4F5F9; padding: 0 2px; border-radius: 2px; }
pre { font-size: 7.5pt; background: #F4F5F9; padding: 6pt; white-space: pre-wrap; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 10pt 0; font-size: 8pt; break-inside: auto; }
th { background: #EEEBFC; text-align: left; }
th, td { border: 1px solid #D5D8E2; padding: 2.5pt 4pt; vertical-align: top; }
tr { break-inside: avoid; }
blockquote { margin: 4pt 0 8pt 0; padding: 4pt 8pt; border-left: 3px solid #5E4AE3; background: #F7F6FE; color: #333; }
blockquote p { margin: 0; }
figure { margin: 6pt 0 10pt 0; text-align: center; break-inside: avoid; }
figure img { max-width: 72%; max-height: 125mm; }
.lado { display: flex; gap: 10pt; align-items: flex-start; break-inside: avoid; }
.lado figure { flex: 1; margin: 4pt 0; }
.lado figure img { max-width: 100%; max-height: 100mm; }
figcaption { font-size: 8.5pt; color: #444; margin-top: 3pt; }
.formula { text-align: center; font-weight: 600; font-size: 10.5pt; margin: 4pt 0 8pt 0; }
.formula p { text-align: center; margin: 0; }
.pendente { background: #FFF1C2; color: #6B4E00; padding: 0 2px; border-radius: 2px; }
ul { margin: 0 0 6pt 0; padding-left: 16pt; } li { margin-bottom: 2pt; }
"""


def gerar_relatorio():
    md = resolver_inclusoes((REL / "RELATORIO_AV1.md").read_text(encoding="utf-8"))
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "rel.md").write_text(md, encoding="utf-8")
        (tmp / "estilo.css").write_text(CSS.replace("FONTE", FONTE.as_uri()), encoding="utf-8")
        html = tmp / "rel.html"
        subprocess.run(["pandoc", str(tmp / "rel.md"), "-f", "markdown+pipe_tables+fenced_divs", "-t", "html5", "-s",
                        "--css", str(tmp / "estilo.css"), "--embed-resources", "--metadata", "lang=pt-BR",
                        "-o", str(html)], check=True)
        texto = html.read_text(encoding="utf-8")
        texto = re.sub(r"⟪([^⟫]*)⟫", r'<span class="pendente">⟪\1⟫</span>', texto)
        html.write_text(texto, encoding="utf-8")
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            nav = p.chromium.launch()
            pg = nav.new_page()
            pg.goto(html.as_uri())
            pg.pdf(path=str(REL / "RELATORIO_AV1.pdf"), format="A4", print_background=True,
                   display_header_footer=True, header_template="<span></span>",
                   footer_template='<div style="font-size:7pt;width:100%;text-align:center;color:#888;">'
                                   'Rastro · AV1 · <span class="pageNumber"></span>/<span class="totalPages"></span></div>',
                   margin={"top": "16mm", "bottom": "16mm", "left": "15mm", "right": "15mm"})
            nav.close()
        subprocess.run(["pandoc", str(tmp / "rel.md"), "-f", "markdown+pipe_tables", "-o", str(REL / "RELATORIO_AV1.docx"),
                        "--resource-path", str(EVID)], check=True)
    print("relatório:", REL / "RELATORIO_AV1.pdf")


def ler_json(nome):
    alvo = RES / nome
    if not alvo.exists():
        return None
    return json.loads(alvo.read_text(encoding="utf-8"))


def gerar_slides():
    import math

    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9, 33,87 × 19,05 cm
    vazio = prs.slide_layouts[6]
    FONTE_NOME = "Plus Jakarta Sans"
    cm = lambda v: Emu(int(v * 360000))
    ICONES = APRES / "icones"
    TOTAL = 11

    # Paleta do aplicativo: violeta domina, verde = certo/detectado, vermelho = erro/limite.
    TINTA, TINTA2, CINZA, CINZA_CLARO = "0D0F14", "1C1F2B", "5F6375", "A9ADBD"
    VIOL, VIOL_CLARO, VIOL_SUAVE = VIOLETA.lstrip("#"), "EEEBFC", "C9C1FF"
    VERDE_, VERDE_CLARO = VERDE.lstrip("#"), "E2F5EC"
    VERM, VERM_CLARO = VERMELHO.lstrip("#"), "FDECEC"
    BRANCO, FUNDO_CARTAO = "FFFFFF", "F4F5F9"
    COR_BLOCO = {"Fundamentação": VIOL, "MVP ao vivo": VERDE_, "Resultados": TINTA}

    # ---------- primitivas ----------
    def forma(sl, x, y, w, h, fill=None, line=None, tipo=MSO_SHAPE.ROUNDED_RECTANGLE, raio_cm=0.3):
        s = sl.shapes.add_shape(tipo, cm(x), cm(y), cm(w), cm(h))
        if tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = min(0.5, raio_cm / min(w, h))
        if fill:
            s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = rgb(line); s.line.width = Pt(1)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def texto(sl, x, y, w, h, paragrafos, tam=16, cor=TINTA, negrito=False, alinh=PP_ALIGN.LEFT,
              ancora=MSO_ANCHOR.TOP, espaco=6, margem=0.0):
        tb = sl.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(margem)
        tf.vertical_anchor = ancora
        if not isinstance(paragrafos, list):
            paragrafos = [paragrafos]
        for i, p in enumerate(paragrafos):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = alinh
            par.space_after = Pt(espaco)
            for parte in (p if isinstance(p, list) else [p]):
                t, o = (parte, {}) if isinstance(parte, str) else parte
                # O "ᵀ" vai num trecho próprio com uma fonte que o tem: na reserva automática, "rᵀ" e "Jᵀ" em
                # negrito saíam como "≡" no PDF do LibreOffice.
                for pedaco in re.split("(ᵀ)", t):
                    if not pedaco:
                        continue
                    r = par.add_run(); r.text = pedaco
                    r.font.size = Pt(o.get("tam", tam)); r.font.bold = o.get("negrito", negrito)
                    r.font.color.rgb = rgb(o.get("cor", cor))
                    r.font.name = "DejaVu Sans" if pedaco == "ᵀ" else FONTE_NOME
        return tb

    def icone(sl, nome, x, y, d, fundo):
        forma(sl, x, y, d, d, fill=fundo, tipo=MSO_SHAPE.OVAL)
        arq = ICONES / f"{nome}.png"
        if arq.exists():
            m = d * 0.24
            sl.shapes.add_picture(str(arq), cm(x + m), cm(y + m), cm(d - 2 * m), cm(d - 2 * m))

    def seta(sl, x1, y1, x2, y2, cor=CINZA):
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cm(x1), cm(y1), cm(x2), cm(y2))
        c.line.color.rgb = rgb(cor); c.line.width = Pt(2.25)
        fim = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
        fim.set("type", "triangle"); fim.set("w", "med"); fim.set("len", "med")

    def base(titulo, bloco, numero, extra="", escuro=False):
        sl = prs.slides.add_slide(vazio)
        fundo = sl.background.fill; fundo.solid(); fundo.fore_color.rgb = rgb(TINTA if escuro else BRANCO)
        cor = VIOL if (escuro and bloco == "Resultados") else COR_BLOCO[bloco]
        larg = 0.95 + 0.235 * len(bloco)
        forma(sl, 1.6, 0.95, larg, 0.72, fill=cor, raio_cm=0.36)
        texto(sl, 1.6, 0.95, larg, 0.72, bloco.upper(), tam=10.5, cor=BRANCO, negrito=True,
              alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
        texto(sl, 1.6, 1.9, 30.6, 1.5, titulo, tam=30, cor=BRANCO if escuro else TINTA, negrito=True, espaco=0)
        cor_rod = CINZA_CLARO if escuro else CINZA
        texto(sl, 1.6, 17.95, 24, 0.6, f"Rastro · Apresentação I · {rodape(bloco, extra, numero)}", tam=10, cor=cor_rod, espaco=0)
        texto(sl, 28.2, 17.95, 4.07, 0.6, f"{numero}/{TOTAL}", tam=10, cor=cor_rod, alinh=PP_ALIGN.RIGHT, espaco=0)
        return sl

    def cartao_formula(sl, x, y, w, h, formulas, legenda, colunas=1, tam_f=20, tam_l=12.5, titulo=None):
        forma(sl, x, y, w, h, fill=TINTA, raio_cm=0.35)
        yy = y + 0.55
        if titulo:
            texto(sl, x + 0.7, yy, w - 1.4, 0.6, titulo, tam=12, cor=VIOL_SUAVE, negrito=True, espaco=0)
            yy += 0.8
        passo = 0.048 * tam_f
        for fml in formulas:
            texto(sl, x + 0.7, yy, w - 1.4, passo, fml, tam=tam_f, cor=BRANCO, negrito=True, espaco=0)
            yy += passo
        yy += 0.35
        por_coluna = -(-len(legenda) // colunas)
        larg_col = (w - 1.4 - 0.5 * (colunas - 1)) / colunas
        for c in range(colunas):
            itens = legenda[c * por_coluna:(c + 1) * por_coluna]
            texto(sl, x + 0.7 + c * (larg_col + 0.5), yy, larg_col, y + h - yy - 0.4,
                  [[(s, {"negrito": True, "cor": VIOL_SUAVE}), ("  " + d, {"cor": "E4E6EE"})] for s, d in itens],
                  tam=tam_l, espaco=3)

    def linha_icone(sl, x, y, w, nome, fundo, titulo, descricao, tam_t=18, tam_d=15, d=1.3):
        icone(sl, nome, x, y, d, fundo)
        texto(sl, x + d + 0.45, y - 0.05, w - d - 0.45, 2.2,
              [[(titulo, {"negrito": True, "tam": tam_t})], [(descricao, {"cor": CINZA, "tam": tam_d})]], espaco=2)

    def estilo_eixo(eixo, tam=12, cor=CINZA):
        eixo.tick_labels.font.size = Pt(tam); eixo.tick_labels.font.name = FONTE_NOME
        eixo.tick_labels.font.color.rgb = rgb(cor)
        eixo.format.line.color.rgb = rgb("D5D8E2")

    def f(v, casas=1, sufixo=""):
        return "—" if v is None else f"{v:.{casas}f}{sufixo}".replace(".", ",")

    def potencia10(v):
        m, e = f"{v:.0e}".split("e")
        return f"{m}·10{str(int(e)).translate(str.maketrans('-0123456789', '⁻⁰¹²³⁴⁵⁶⁷⁸⁹'))}"

    # ---------- dados (nenhum número medido é digitado aqui) ----------
    est, aq, rec = ler_json("estatico.json"), ler_json("aquisicao.json"), ler_json("reconciliacao.json")
    aq_antes = ler_json("aquisicao_antes_d23.json")
    P = "⟪após a coleta⟫"

    def resumo_testes():
        alvo = RES / "testes_unidade.md"
        if not alvo.exists():
            return None
        t = alvo.read_text(encoding="utf-8")
        ls = re.findall(r"^\|\s*(com\.[^|]+)\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|", t, flags=re.M)
        c = re.search(r"Commit: `([0-9a-f]+)`", t)
        return (sum(int(l[1]) for l in ls), len(ls), sum(int(l[2]) + int(l[3]) for l in ls), c.group(1) if c else None)

    def resumo_corrida():
        alvo = RES / "corrida_10_execucoes.md"
        if not alvo.exists():
            return None
        r = re.findall(r"^\|\s*\d+\s*\|.*\|\s*(PASSOU|FALHOU)\s*\|\s*$", alvo.read_text(encoding="utf-8"), flags=re.M)
        return (r.count("PASSOU"), len(r)) if r else None

    def ler_sensibilidade():
        alvo = RES / "sensibilidade.md"
        if not alvo.exists():
            return None
        linhas = re.findall(r"^\|\s*([^|]+?)\s*\|\s*([\d.]+)\s*\|\s*\d+\s*\|\s*\d+\s*\|\s*\d+\s*\|\s*(—|\d+%)\s*\|"
                            r"\s*(?:—|\d+%)\s*\|\s*(\d+)%\s*\|", alvo.read_text(encoding="utf-8"), flags=re.M)
        tab = {}
        for deg, par, det, fa in linhas:
            tab.setdefault(deg, []).append((float(par), None if det == "—" else int(det[:-1]), int(fa)))
        return tab

    tst, cor, sens = resumo_testes(), resumo_corrida(), ler_sensibilidade()
    semente = rec["semente"] if rec else 20261020
    salto_min = rec.get("salto_minimo_90pct_m") if rec else None

    # ---------- 1 — título ----------
    sl = prs.slides.add_slide(vazio)
    fundo = sl.background.fill; fundo.solid(); fundo.fore_color.rgb = rgb(TINTA)
    logo = RAIZ / "rastro-logo-kit/logo/rastro-simbolo-branco-1024.png"
    if logo.exists():
        sl.shapes.add_picture(str(logo), cm(1.6), cm(1.5), height=cm(2.2))
    texto(sl, 1.6, 4.9, 30, 2.4, "Rastro", tam=60, cor=BRANCO, negrito=True, espaco=0)
    texto(sl, 1.6, 7.35, 30, 1.4, "sente, corrige e detecta", tam=30, cor=VIOL_SUAVE, espaco=0)
    texto(sl, 1.6, 9.0, 30, 0.9, "Apresentação I · AV1 · GAT108 Automação Avançada · UFLA · 27/10/2026",
          tam=16, cor=CINZA_CLARO, espaco=0)
    texto(sl, 1.6, 10.5, 30, 0.9, "Juliano Cesar Ferreira Ramos  ·  Andre Bueno Ocanha  ·  João Paulo Silva",
          tam=18, cor=BRANCO, espaco=0)
    for i, (bloco, minutos, fundo_c) in enumerate([("Fundamentação", "8 min", VIOL), ("MVP ao vivo", "9 min", VERDE_),
                                                    ("Resultados", "5 min", TINTA2)]):
        x = 1.6 + i * 10.35
        forma(sl, x, 13.9, 9.85, 2.0, fill=fundo_c, raio_cm=0.4)
        linhas = [[(f"{bloco} · {minutos}", {"negrito": True, "tam": 16})]]
        if APRESENTADORES.get(bloco):
            linhas.append([(APRESENTADORES[bloco], {"tam": 13, "cor": "E4E6EE"})])
        texto(sl, x + 0.6, 13.9, 8.7, 2.0, linhas, cor=BRANCO, ancora=MSO_ANCHOR.MIDDLE, espaco=2)
    texto(sl, 1.6, 16.5, 30, 0.7, "22 min de exposição e 8 min de arguição", tam=12, cor=CINZA_CLARO, espaco=0)

    # ---------- 2 — o que a camada 1 mede ----------
    sl = base("O que a camada 1 mede", "Fundamentação", 2)
    texto(sl, 1.6, 3.9, 30.6, 1.4, [[("Rede mesh entre Androids em áreas sem cobertura. Na AV1, o nó precisa ", {}),
                                      ("sentir, corrigir e detectar", {"negrito": True, "cor": VIOL}),
                                      (" sem enviar nada para fora.", {})]], tam=18)
    linha_icone(sl, 1.6, 6.0, 16.5, "satelite", VIOL, "Fixo GNSS a cada 1 s",
                "posição, velocidade (Doppler) e rumo, cada um com a acurácia de 68% que o receptor informa")
    linha_icone(sl, 1.6, 9.45, 16.5, "atividade", VIOL, "Acelerômetro a 50 Hz",
                "aquisição contínua no serviço em primeiro plano, também com a tela bloqueada")
    linha_icone(sl, 1.6, 12.9, 16.5, "incerteza", VIOL, "De onde vem a incerteza",
                "multipercurso, geometria dos satélites e céu obstruído")
    forma(sl, 19.6, 5.9, 12.67, 11.1, fill=VIOL_CLARO, raio_cm=0.4)
    icone(sl, "registro", 20.3, 6.55, 1.1, VIOL)
    texto(sl, 21.7, 6.6, 10, 1.0, "RegistroFix", tam=20, negrito=True, cor=TINTA, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    campos = [("tempo", "monotônico (ns) e UTC"), ("lat, lon", "graus"), ("r68", "raio horizontal de 68% (m)"),
              ("s ± σs", "velocidade (m/s)"), ("θ ± σθ", "rumo a partir do norte (°)"), ("provedor", "gps ou network")]
    texto(sl, 20.3, 8.2, 11.3, 6.0, [[(c, {"negrito": True, "cor": VIOL}), ("   " + d, {"cor": TINTA})] for c, d in campos],
          tam=16, espaco=8)
    texto(sl, 20.3, 14.75, 11.3, 1.9, "Campo ausente fica vazio: nenhuma incerteza é inventada.",
          tam=16, negrito=True, cor=VIOL)

    # ---------- 3, 4 e 7: desenho de Juliano Cesar Ferreira Ramos (commit 28c18ec), na grade deste deck ----------
    # A grade dele tem margem de 1,2 cm e conteúdo de 3,0 a 16,95 cm; J() leva cada caixa para esta grade.
    def J(x, y, w, h):
        kx, ky = 30.67 / 31.47, 13.2 / 13.95
        return 1.6 + (x - 1.2) * kx, 3.9 + (y - 3.0) * ky, w * kx, h * ky

    def caixaJ(sl, x, y, w, h, fill, raio_cm=0.35, tipo=MSO_SHAPE.ROUNDED_RECTANGLE):
        return forma(sl, *J(x, y, w, h), fill=fill, raio_cm=raio_cm, tipo=tipo)

    def textoJ(sl, x, y, w, h, paragrafos, **kw):
        return texto(sl, *J(x, y, w, h), paragrafos, **kw)

    def selo(sl, x, y, n):
        caixaJ(sl, x, y, 1.6, 1.8, VIOL, raio_cm=0.2)
        textoJ(sl, x, y, 1.6, 1.8, str(n), tam=18, negrito=True, cor=BRANCO, alinh=PP_ALIGN.CENTER,
               ancora=MSO_ANCHOR.MIDDLE, espaco=0)

    def oh(v):
        return (v, {"negrito": True, "cor": VIOL})

    # 3 — estruturas
    sl = base("Memória estável durante a coleta", "Fundamentação", 3)
    textoJ(sl, 1.2, 3.0, 31.47, 1.28, [[("Duas estruturas, uma regra: ", {"negrito": True}),
                                         ("a memória não cresce.", {"negrito": True, "cor": VIOL})]],
           tam=20, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    estruturas = [
        (1.2, 1, "Trajeto", [("guarda somente as últimas posições.", {})],
         [("tamanho ", {}), oh("O(1)"), (" · média e recorte ", {}), oh("O(n)"), (" · inserção ", {}), oh("O(n)")]),
        (17.43, 2, "Buffer IMU", [("guarda ", {}), ("500 leituras", {"negrito": True, "cor": VIOL}),
                                  (", cerca de 10 segundos.", {})],
         [("inserção ", {}), oh("O(1)"), (" · descarte ", {}), oh("O(1)"), (" · cópia na leitura ", {}), oh("O(n)")]),
    ]
    for x, n, nome, desc, compl in estruturas:
        caixaJ(sl, x, 4.83, 15.23, 3.94, FUNDO_CARTAO)
        selo(sl, x + 0.67, 5.91, n)
        textoJ(sl, x + 2.82, 4.83, 11.97, 3.94, [[(nome, {"negrito": True, "tam": 20})], desc,
                                                  [(a, dict(o, tam=12.5, cor=o.get("cor", CINZA))) for a, o in compl]],
               tam=15, ancora=MSO_ANCHOR.MIDDLE, espaco=4)
    caixaJ(sl, 1.2, 9.44, 31.47, 4.0, FUNDO_CARTAO)
    textoJ(sl, 1.87, 9.44, 11.67, 4.0, [[("Chegou uma leitura e está cheio?", {"negrito": True})],
                                         [("A mais antiga sai.", {"negrito": True, "cor": VIOL})]],
           tam=19, ancora=MSO_ANCHOR.MIDDLE, espaco=2)
    for i, c in enumerate(["D5D8E3"] + ["B9B0F4"] * 5 + [VIOL]):
        caixaJ(sl, 14.31 + i * 2.47, 10.11, 2.11, 1.78, c, raio_cm=0.25)
    textoJ(sl, 13.83, 12.06, 3.06, 0.94, "←", tam=16, cor=CINZA, alinh=PP_ALIGN.CENTER, espaco=0)
    textoJ(sl, 16.78, 12.06, 12.0, 0.94, "500 leituras · ≈ 10 s", tam=13, negrito=True, alinh=PP_ALIGN.CENTER, espaco=0)
    textoJ(sl, 28.67, 12.06, 3.06, 0.94, "→", tam=16, cor=VIOL, alinh=PP_ALIGN.CENTER, espaco=0)
    caixaJ(sl, 1.2, 14.06, 15.23, 2.89, FUNDO_CARTAO)
    textoJ(sl, 1.87, 14.06, 13.9, 2.89, [[("Quem lê o buffer recebe ", {}), ("uma cópia", {"negrito": True, "cor": VIOL}),
                                          (": a leitura não se mistura com uma gravação em andamento.", {})]],
           tam=15, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    caixaJ(sl, 17.43, 14.06, 15.23, 2.89, VIOL)
    textoJ(sl, 18.1, 14.06, 13.9, 2.89, "Assim, o app pode ficar ligado por muito tempo sem aumentar o uso de memória.",
           tam=15, negrito=True, cor=BRANCO, ancora=MSO_ANCHOR.MIDDLE, espaco=0)

    # 4 — incerteza
    sl = base("A precisão vem do próprio GPS", "Fundamentação", 4)
    passos4 = [[("O GPS não entrega só a posição: ele também informa uma ", {}), ("margem de erro", {"negrito": True}), (".", {})],
               [("O Android representa essa margem como um ", {}), ("círculo de precisão de 68%", {"negrito": True}), (".", {})],
               [("Nós transformamos esse círculo em uma ", {}), ("medida de incerteza", {"negrito": True, "cor": VIOL}),
                (" para a correção.", {})]]
    for i, p in enumerate(passos4):
        y = 3.0 + i * 2.11
        selo(sl, 1.2, y, i + 1)
        textoJ(sl, 3.2, y, 14.67, 1.8, [p], tam=15, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    exemplo = 6.0
    caixaJ(sl, 1.2, 9.44, 16.67, 2.61, VIOL)
    textoJ(sl, 1.87, 9.44, 5.3, 2.61, [[(f(exemplo, 0, " m"), {"negrito": True, "tam": 26})],
                                        [("precisão r68 (exemplo)", {"tam": 11.5, "cor": "E4E0FF"})]],
           cor=BRANCO, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    textoJ(sl, 7.05, 9.44, 3.3, 2.61, [[("→", {"negrito": True, "tam": 22})], [("σ = r68 / √2,279", {"tam": 11.5, "cor": "E4E0FF"})]],
           cor=BRANCO, alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    textoJ(sl, 10.6, 9.44, 7.0, 2.61, [[(f"≈ {f(exemplo / (-2 * math.log(0.32)) ** 0.5, 0, ' m')}", {"negrito": True, "tam": 26})],
                                        [("de incerteza σ em cada direção", {"tam": 11.5, "cor": "E4E0FF"})]],
           cor=BRANCO, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    caixaJ(sl, 1.2, 12.56, 16.67, 4.39, FUNDO_CARTAO)
    textoJ(sl, 1.87, 12.56, 15.33, 4.39, [
        [("Com a velocidade é igual: usamos o valor, a direção e as margens informadas pelo sensor.", {})],
        [("Se alguma margem não vier informada, o aplicativo ", {}), ("não inventa um número", {"negrito": True, "cor": VIOL}),
         (". Assim, uma leitura incerta não é tratada como se fosse confiável.", {})]],
        tam=14, ancora=MSO_ANCHOR.MIDDLE, espaco=8)
    caixaJ(sl, 18.78, 3.0, 13.89, 13.94, FUNDO_CARTAO)
    svg = RES / "dispersao_estatica.svg"
    png = None
    if svg.exists():
        destino = Path(tempfile.mkdtemp()) / "dispersao_estatica.png"
        if shutil.which("rsvg-convert"):
            subprocess.run(["rsvg-convert", "-w", "1400", "-o", str(destino), str(svg)], check=False)
        else:
            subprocess.run(["soffice", "--headless", "--convert-to", "png", "--outdir", str(destino.parent), str(svg)],
                           check=False, capture_output=True)
        png = destino if destino.exists() else None
    if png:
        x, y, w, h = J(19.81, 3.5, 11.83, 12.28)
        alt = h
        larg = alt * 560 / 600
        sl.shapes.add_picture(str(png), cm(x + (w - larg) / 2), cm(y), cm(larg), cm(alt))
    textoJ(sl, 19.11, 15.78, 13.22, 0.94, "Tracejado: precisão de 68% declarada pelo GPS", tam=11.5, cor=CINZA,
           alinh=PP_ALIGN.CENTER, espaco=0)

    # ---------- 5 — reconciliação ----------
    sl = base("Reconciliação sob a restrição cinemática", "Fundamentação", 5)
    cartao_formula(sl, 1.6, 3.9, 30.67, 4.6, ["pₐ(k+1) − pₐ(k) − ½·(vₐ(k) + vₐ(k+1))·Δtₖ = 0"], [
        ("pₐ(k)", "posição do fixo k no eixo a (m, projeção local)"), ("vₐ(k)", "componente da velocidade (m/s)"),
        ("Δtₖ", "intervalo entre os fixos (s, relógio monotônico)"), ("a", "eixo leste ou norte"),
        ("k", "fixo da janela, de 1 a N = 10"), ("linear", "Δt é conhecido → A·y = 0")],
        colunas=2, tam_f=22, titulo="Restrição física: integração trapezoidal da velocidade")
    cartao_formula(sl, 1.6, 8.95, 17.6, 6.0, ["ŷ = y − Σ·Aᵀ·(A·Σ·Aᵀ)⁻¹·A·y", "γ = rᵀ·V⁻¹·r"], [
        ("y, ŷ", "medidas e reconciliadas"), ("Σ", "covariância do sensor"),
        ("A", "2(N−1) restrições × 4N medidas"), ("r", "resíduo A·y"), ("V", "covariância A·Σ·Aᵀ"),
        ("γ", "χ² com 2(N−1) graus de liberdade")], colunas=2, tam_f=19, titulo="Mínimos quadrados ponderados")
    forma(sl, 19.7, 8.95, 12.57, 6.0, fill=VERDE_CLARO, raio_cm=0.4)
    icone(sl, "alvo", 20.35, 9.55, 1.2, VERDE_)
    texto(sl, 20.35, 11.05, 11.3, 3.8, [[("A medida mais incerta é a que mais se move.", {"negrito": True, "tam": 17})],
                                         [("A correção y − ŷ é proporcional a Σ. No teste com dois fixos de r68 = 2 m e "
                                           "20 m, o segundo absorve cerca de 99% da correção.", {"cor": CINZA, "tam": 13.5})]],
          espaco=5)
    descartes = [("Mínimos quadrados sem peso", "puxa igual a medida boa e a ruim"),
                 ("Filtro de Kalman", "estima estado e exige ruído Q arbitrado"),
                 ("|Δp| = s·Δt", "não linear e perde a direção")]
    texto(sl, 1.6, 15.35, 4.0, 1.6, "Descartadas", tam=14, negrito=True, cor=VERM, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    for i, (n, motivo) in enumerate(descartes):
        x = 5.4 + i * 9.0
        forma(sl, x, 15.35, 8.6, 1.65, fill=VERM_CLARO, raio_cm=0.3)
        texto(sl, x + 0.4, 15.35, 7.9, 1.65, [[(n, {"negrito": True, "tam": 13})], [(motivo, {"cor": CINZA, "tam": 11.5})]],
              ancora=MSO_ANCHOR.MIDDLE, espaco=0)

    # ---------- 6 — erro grosseiro ----------
    sl = base("Erro grosseiro: critério quantitativo de corte", "Fundamentação", 6)
    cartao_formula(sl, 1.6, 3.9, 16.6, 10.6, ["zⱼ = |aⱼ| / √Wⱼⱼ", "β = 1 − (1 − α)^(1/n)", "z_crítico = Φ⁻¹(1 − β/2) ≈ 3,2"], [
        ("aⱼ", "ajuste da medida j: yⱼ − ŷⱼ"), ("Wⱼⱼ", "variância do ajuste, W = Σ·Aᵀ·V⁻¹·A·Σ"),
        ("zⱼ", "resíduo normalizado: N(0,1) sem erro grosseiro"), ("α", "5% de falso alarme por janela"),
        ("n", "40 medidas por janela (1,96 em cada daria ~87%)"), ("β", "nível de cada teste (correção de Šidák)"),
        ("Φ⁻¹", "quantil da normal padrão"), ("corte", "retira o fixo e reconcilia de novo, até 3 vezes")],
        tam_f=20, titulo="Teste de medida com eliminação serial")
    texto(sl, 19.1, 3.9, 13.2, 1.5, [[("Sem o detector, a reconciliação piora", {"negrito": True, "tam": 17})],
                                       [("salto de 30 m no fixo central · afastamento RMS dos vizinhos", {"cor": CINZA, "tam": 12.5})]],
          espaco=2)
    if rec:
        cd = CategoryChartData()
        cd.categories = ["Sem detector", "Com detector"]
        cd.add_series("RMS dos vizinhos (m)", (rec["demo_rms_sem_detector_m"], rec["demo_rms_com_detector_m"]))
        gf = sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, cm(18.9), cm(5.6), cm(13.4), cm(6.6), cd)
        ch = gf.chart
        ch.has_legend = False; ch.has_title = False
        ch.font.name = FONTE_NOME; ch.font.size = Pt(13)
        plot = ch.plots[0]; plot.gap_width = 55; plot.vary_by_categories = False
        ser = plot.series[0]
        for i, (cor_b, rot) in enumerate([(VERM, f(rec["demo_rms_sem_detector_m"], 2, " m")),
                                           (VERDE_, f(rec["demo_rms_com_detector_m"], 2, " m"))]):
            pt = ser.points[i]
            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(cor_b)
            dl = pt.data_label; dl.has_text_frame = True; dl.text_frame.text = rot
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            for r in dl.text_frame.paragraphs[0].runs:
                r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = rgb(cor_b); r.font.name = FONTE_NOME
        va = ch.value_axis
        va.visible = False; va.has_major_gridlines = False
        va.minimum_scale = 0; va.maximum_scale = rec["demo_rms_sem_detector_m"] * 1.35
        ca = ch.category_axis
        ca.reverse_order = True; estilo_eixo(ca, 14, TINTA); ca.format.line.fill.background()
    texto(sl, 19.1, 12.35, 13.2, 2.2, (f"Referência: a mesma janela reconciliada sem o salto. O fixo corrompido foi localizado "
                                        f"em {rec['demo_localizou']} de {rec['janelas']} janelas reais.") if rec else P,
          tam=12.5, cor=CINZA)
    forma(sl, 1.6, 15.0, 30.67, 2.0, fill=FUNDO_CARTAO, raio_cm=0.35)
    icone(sl, "alerta", 2.15, 15.45, 1.1, VERM)
    texto(sl, 3.65, 15.0, 28.2, 2.0, [[("Limite: ", {"negrito": True}),
                                        ("nos fixos das pontas há uma restrição por eixo e todas as medidas dela têm o mesmo z. "
                                         "O erro é detectado, mas não localizado: por isso a reprodução usa janela deslizante.", {})]],
          tam=14, ancora=MSO_ANCHOR.MIDDLE, espaco=0)

    # 7 — concorrência e aquisição (desenho de Juliano Cesar Ferreira Ramos)
    sl = base("Coleta contínua, sem conflito", "Fundamentação", 7)
    for i, (t7, d7) in enumerate([("Acelerômetro", "50 leituras por segundo"), ("GPS", "um fixo por segundo"),
                                   ("Consumidor", "organiza os dados"), ("Registro", "salva os dados")]):
        y = 3.0 + i * 2.39
        caixaJ(sl, 1.2, y, 7.5, 2.0, FUNDO_CARTAO)
        textoJ(sl, 1.7, y, 6.78, 2.0, [[(t7, {"negrito": True, "tam": 15})], [(d7, {"tam": 12, "cor": CINZA})]],
               ancora=MSO_ANCHOR.MIDDLE, espaco=0)
        caixaJ(sl, 8.97, y + 0.64, 1.06, 0.72, "B9B0F4", tipo=MSO_SHAPE.RIGHT_ARROW)
    caixaJ(sl, 10.31, 3.0, 7.36, 9.17, VIOL)
    textoJ(sl, 10.6, 3.0, 6.8, 7.06, [[("Buffer e gravador", {"negrito": True, "tam": 22})],
                                       [("Cada um protegido por um único cadeado.", {"tam": 13, "cor": "E4E0FF"})]],
           cor=BRANCO, alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE, espaco=6)
    caixaJ(sl, 11.21, 10.39, 5.56, 1.17, BRANCO, raio_cm=0.55)
    textoJ(sl, 11.21, 10.39, 5.56, 1.17, "synchronized", tam=13, negrito=True, cor=VIOL, alinh=PP_ALIGN.CENTER,
           ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    caixaJ(sl, 1.2, 12.83, 16.47, 4.11, FUNDO_CARTAO)
    textoJ(sl, 1.87, 12.83, 15.13, 4.11, [
        [("synchronized", {"negrito": True, "cor": VIOL}), (" funciona como um cadeado:", {"negrito": True})],
        [("apenas uma parte mexe em cada estrutura por vez.", {})],
        [("Descartados: ConcurrentLinkedQueue (não limita nem conta o descarte) e ReentrantLock (sem uso para tryLock). "
          "Escalonamento: CFS do Android, sem tempo real.", {"tam": 11, "cor": CINZA})]],
        tam=15, ancora=MSO_ANCHOR.MIDDLE, espaco=4)

    def numero7(x, y, v, rot):
        textoJ(sl, x, y, 6.11, 3.06, [[(v, {"negrito": True, "tam": 26, "cor": VIOL})], [(rot, {"tam": 12, "cor": CINZA})]], espaco=0)
    caixaJ(sl, 18.5, 3.0, 14.17, 5.39, FUNDO_CARTAO)
    textoJ(sl, 19.17, 3.5, 12.83, 1.11, "Resultado no aparelho", tam=15, negrito=True, espaco=0)
    numero7(19.17, 4.94, f(aq["bloqueado_min"], 1, " min") if aq else P, "com a tela bloqueada")
    numero7(25.89, 4.94, f(aq["imu_hz_medio"], 1) if aq else P, "leituras por segundo na IMU")
    caixaJ(sl, 18.5, 9.06, 14.17, 7.89, FUNDO_CARTAO)
    textoJ(sl, 19.17, 9.56, 12.83, 1.11, "Teste de corrida", tam=15, negrito=True, espaco=0)
    numero7(19.17, 11.0, f"{cor[1]}" if cor else P, "execuções independentes")
    numero7(25.89, 11.0, "10 mil", "leituras por rodada, 10 rodadas cada")
    caixaJ(sl, 19.17, 14.56, 12.83, 1.72, VIOL)
    textoJ(sl, 19.17, 14.56, 12.83, 1.72,
           "Nenhuma inconsistência observada." if cor and cor[0] == cor[1] else (f"{cor[1] - cor[0]} execuções com falha" if cor else P),
           tam=14, negrito=True, cor=BRANCO, alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE, espaco=0)

    # ---------- 8 — ensaio reprodutível ----------
    sl = base("Ensaio reprodutível e injeção de falhas", "Fundamentação", 8)
    forma(sl, 1.6, 3.9, 13.6, 3.0, fill=VIOL_CLARO, raio_cm=0.4)
    icone(sl, "semente", 2.2, 4.45, 1.3, VIOL)
    texto(sl, 3.9, 4.15, 11.0, 2.6, [[(f"Semente {semente}", {"negrito": True, "tam": 18})],
                                      [("java.util.Random tem algoritmo especificado: o mesmo resultado em qualquer JVM ou Android. "
                                        "Um gerador por (semente, rótulo).", {"cor": CINZA, "tam": 12.5})]], espaco=2)
    linha_icone(sl, 1.6, 7.55, 13.6, "ruido", VIOL, "Ruído",
                "gaussiano extra, sem mudar a acurácia declarada (sensor otimista)", tam_t=16, tam_d=13, d=1.15)
    linha_icone(sl, 1.6, 10.45, 13.6, "perda", VIOL, "Perda",
                "fixos retirados com probabilidade p (a janela encurta e Δt cresce)", tam_t=16, tam_d=13, d=1.15)
    linha_icone(sl, 1.6, 13.35, 13.6, "salto", VERM, "Salto",
                "um fixo deslocado de J metros em direção sorteada: o erro grosseiro", tam_t=16, tam_d=13, d=1.15)
    texto(sl, 16.4, 3.9, 15.9, 1.0, "Onde a detecção começa a falhar", tam=17, negrito=True, espaco=0)
    saltos = [s for s in (sens or {}).get("salto", []) if s[1] is not None]
    if saltos:
        cd = CategoryChartData()
        cd.categories = [f"{s[0]:g}" for s in saltos]
        cd.add_series("Detecção do salto", [s[1] for s in saltos])
        cd.add_series("Falso alarme", [s[2] for s in saltos])
        cd.add_series("90%", [90] * len(saltos))
        gf = sl.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, cm(16.1), cm(4.9), cm(16.2), cm(8.6), cd)
        ch = gf.chart
        ch.font.name = FONTE_NOME; ch.font.size = Pt(12)
        ch.has_title = False
        ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(12); ch.legend.font.color.rgb = rgb(CINZA)
        estilos = [(VIOL, 2.75, None, XL_MARKER_STYLE.CIRCLE), (VERM, 1.75, None, XL_MARKER_STYLE.CIRCLE),
                   ("A9ADBD", 1.5, MSO_LINE_DASH_STYLE.DASH, XL_MARKER_STYLE.NONE)]
        for ser, (c, larg, traco, marca) in zip(ch.plots[0].series, estilos):
            ser.smooth = False
            ser.format.line.color.rgb = rgb(c); ser.format.line.width = Pt(larg)
            if traco:
                ser.format.line.dash_style = traco
            ser.marker.style = marca
            if marca != XL_MARKER_STYLE.NONE:
                ser.marker.size = 7
                ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(c)
                ser.marker.format.line.color.rgb = rgb(c)
        det = ch.plots[0].series[0]
        for i, s in enumerate(saltos):
            if s[0] in (2.0, 5.0):
                dl = det.points[i].data_label; dl.has_text_frame = True; dl.text_frame.text = f"{s[1]}%"
                dl.position = XL_LABEL_POSITION.ABOVE
                for r in dl.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = rgb(VIOL); r.font.name = FONTE_NOME
        va = ch.value_axis
        va.minimum_scale = 0; va.maximum_scale = 100; va.major_unit = 25
        va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb("E6E8EF")
        va.tick_labels.number_format = '0"%"'; va.tick_labels.number_format_is_linked = False
        estilo_eixo(va); va.format.line.fill.background()
        ca = ch.category_axis; estilo_eixo(ca)
        ca.has_title = True; ca.axis_title.text_frame.text = "salto injetado (m)"
        for r in ca.axis_title.text_frame.paragraphs[0].runs:
            r.font.size = Pt(12); r.font.bold = False; r.font.color.rgb = rgb(CINZA); r.font.name = FONTE_NOME
    ruido_1m = next((fa for par, _, fa in (sens or {}).get("ruído σ (m)", []) if par == 1.0), None)
    perdas_ok = [par for par, d, _ in (sens or {}).get("perda (prob.)", []) if d == 100]
    destaques = [(f(salto_min, 0, " m") if salto_min is not None else P, "menor salto com ≥ 90% de detecção", VIOL),
                 (f"{ruido_1m}%" if ruido_1m is not None else P, "de falso alarme com 1 m de ruído não declarado", VERM),
                 (f"{max(perdas_ok) * 100:.0f}%" if perdas_ok else P, "de perda de fixos sem mudar a detecção", VERDE_)]
    for i, (v, rot, c) in enumerate(destaques):
        x = 16.4 + i * 5.37
        forma(sl, x, 13.85, 5.0, 3.15, fill=FUNDO_CARTAO, raio_cm=0.35)
        texto(sl, x + 0.4, 14.0, 4.3, 2.9, [[(v, {"negrito": True, "tam": 22, "cor": c})], [(rot, {"tam": 11, "cor": CINZA})]],
              espaco=0)

    # ---------- 9 — MVP ----------
    sl = base("MVP ao vivo: o caminho do dado", "MVP ao vivo", 9, extra="9 min")
    passos = [("celular", "Aparelho bloqueado", "notificação com fixos GNSS, IMU e tempo subindo"),
              ("lupa", "Tela Ensaio AV1", "toque longo no status: aquisição ao vivo"),
              ("play", "Reprodução ×10", f"caminhada real, salto de 30 m, semente {semente}"),
              ("grafico", "Resíduo na tela", "RMS antes → depois e γ contra o limite χ² de 95%"),
              ("alvo", "Erro grosseiro", "“ERRO GROSSEIRO no fixo N ← o salto injetado”"),
              ("alerta", "Falha de propósito",
               f"salto de 2 m, abaixo do mínimo detectável ({f(salto_min, 0, ' m')}): passa sem ser detectado"
               if salto_min is not None else "salto abaixo do mínimo detectável: passa sem ser detectado")]
    for i, (ic, t, d) in enumerate(passos):
        col, lin = i % 3, i // 3
        x, y = 1.6 + col * 10.35, 4.0 + lin * 5.55
        falha = i == 5
        forma(sl, x, y, 9.97, 5.0, fill=VERM_CLARO if falha else FUNDO_CARTAO, raio_cm=0.4)
        icone(sl, ic, x + 0.6, y + 0.6, 1.3, VERM if falha else VERDE_)
        texto(sl, x + 7.6, y + 0.45, 1.8, 1.4, f"{i + 1}", tam=28, negrito=True, cor=VERM if falha else "C4C8D4",
              alinh=PP_ALIGN.RIGHT, espaco=0)
        texto(sl, x + 0.6, y + 2.25, 8.8, 2.6, [[(t, {"negrito": True, "tam": 17})], [(d, {"cor": CINZA, "tam": 13})]], espaco=3)
    forma(sl, 1.6, 15.2, 30.67, 1.8, fill=VERDE_CLARO, raio_cm=0.35)
    texto(sl, 2.2, 15.2, 29.5, 1.8, [[("Antes da aula: ", {"negrito": True}),
                                       ("app ativado 30+ min antes com a tela bloqueada · localização “Permitir o tempo todo” · "
                                        "economia de energia desligada", {})]], tam=13.5, ancora=MSO_ANCHOR.MIDDLE, espaco=0)

    # ---------- 10 — resultados ----------
    sl = base("Resultados com a condição de medição", "Resultados", 10, extra="5 min")
    cartoes = [
        (f"{tst[0]}" if tst else P, "testes de unidade", f"{tst[2]} falhas · {tst[1]} classes · gradlew test no {tst[3]}" if tst else "", "teste"),
        (f"{cor[0]} de {cor[1]}" if cor else P, "execuções do teste de corrida", "sem falha · execuções independentes", "anel"),
        (f(aq["bloqueado_min"], 1, " min") if aq else P, "de aquisição com a tela bloqueada",
         f"IMU {f(aq['imu_hz_medio'], 1, ' Hz')} · {aq['imu_descartadas']} leituras perdidas" if aq else "", "relogio"),
        (f(aq["cobertura_gps"] * 100, 0, "%") if aq else P, "do tempo com GPS, tela apagada",
         (f"{f(aq_antes['cobertura_gps'] * 100, 0, '%')} antes da D-23 · " if aq_antes else "")
         + (f"rede a cada ≤ {f(aq['maior_intervalo_qualquer_fixo_s'], 0, ' s')}" if aq else ""), "satelite"),
        (f(aq["consumo_pct_h"], 1, " %/h") if aq else P, "de bateria", "no mesmo trecho bloqueado", "bateria"),
        (f(est["razao_otimismo"], 2) if est else P, "razão de otimismo",
         f"parado · {f(est['fracao_repetidos'] * 100, 0, '%')} dos fixos retidos: não mede o ruído" if est else "", "regua"),
        (f"{f(rec['rms_antes_medio_m'], 2, ' m')} → {potencia10(rec['rms_depois_max_m'])} m" if rec else P,
         "resíduo RMS médio, antes → depois", f"{rec['janelas']} janelas reais da caminhada" if rec else "", "formula"),
        (f(salto_min, 0, " m") if salto_min is not None else P, "menor salto detectável",
         f"≥ 90% de detecção · semente {semente}", "alvo"),
    ]
    for i, (v, rot, cond, ic) in enumerate(cartoes):
        col, lin = i % 4, i // 4
        x, y = 1.6 + col * 7.75, 3.95 + lin * 6.15
        forma(sl, x, y, 7.42, 5.75, fill=FUNDO_CARTAO, raio_cm=0.4)
        icone(sl, ic, x + 0.5, y + 0.5, 1.05, TINTA)
        longo = len(v) > 9
        texto(sl, x + 0.5, y + 1.85, 6.6, 1.5, v, tam=20 if longo else 30, negrito=True, cor=VIOL,
              ancora=MSO_ANCHOR.MIDDLE, espaco=0)
        texto(sl, x + 0.5, y + 3.4, 6.6, 2.3, [[(rot, {"negrito": True, "tam": 13.5})], [(cond, {"cor": CINZA, "tam": 11})]],
              espaco=2)
    texto(sl, 1.6, 16.35, 30.67, 1.2, "Samsung SM-S926B · coletas de 22/10/2026 no campus da UFLA (sessão bloqueada de 21h12 "
          "fora do campus, coordenadas deslocadas) · scripts\\av1_evidencias.ps1 · docs/evidencias/resultados",
          tam=10.5, cor=CINZA)

    # ---------- 11 — limitações ----------
    sl = base("Limitações e próximo passo", "Resultados", 11, escuro=True)
    gps = (f"Com a tela apagada, GPS em {f(aq['cobertura_gps'] * 100, 0, '%')} do tempo e uma pausa de "
           f"{f(aq['maior_intervalo_fixos_s'] / 60, 1, ' min')} sem causa conhecida.") if aq else P
    colunas = [("O que ainda não sabemos", "pergunta", VERM, [
                    "Sem verdade de campo: a referência é a janela limpa reconciliada.",
                    "O erro do GNSS é correlacionado no tempo: γ não é exatamente χ².", gps,
                    "Parado, o receptor retém a posição: o ensaio estático não mede o ruído.",
                    "O registro fica em claro até a cifragem da AV2."]),
               ("AV2", "foguete", VIOL, [
                    "Pedir a localização em segundo plano no app; consumidor IMU com taxa fixa (D-22).",
                    "Cᵢ medido e escalonabilidade das quatro tarefas periódicas.",
                    "Aprendizado online dos encontros de rede.",
                    "Agente com RuleEngine e modelo removível.",
                    "AES-GCM com chave no Keystore e frota com deriva de relógio."])]
    for i, (t, ic, c, itens) in enumerate(colunas):
        x = 1.6 + i * 15.6
        forma(sl, x, 3.9, 15.07, 13.1, fill=TINTA2, raio_cm=0.4)
        icone(sl, ic, x + 0.6, 4.5, 1.2, c)
        texto(sl, x + 2.1, 4.5, 12.5, 1.2, t, tam=20, negrito=True, cor=BRANCO, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
        for j, it in enumerate(itens):
            yy = 6.35 + j * 2.1
            forma(sl, x + 0.75, yy + 0.22, 0.32, 0.32, fill=c, tipo=MSO_SHAPE.OVAL)
            texto(sl, x + 1.45, yy, 13.1, 2.0, it, tam=14.5, cor="E4E6EE", espaco=0)

    APRES.mkdir(parents=True, exist_ok=True)
    pptx = APRES / "apresentacao_I_AV1.pptx"
    prs.save(str(pptx))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(APRES), str(pptx)], check=True,
                   capture_output=True)
    print("slides:", pptx, "e", pptx.with_suffix(".pdf"))


if __name__ == "__main__":
    alvo = sys.argv[1] if len(sys.argv) > 1 else "tudo"
    if alvo in ("tudo", "relatorio"):
        gerar_relatorio()
    if alvo in ("tudo", "slides"):
        gerar_slides()
